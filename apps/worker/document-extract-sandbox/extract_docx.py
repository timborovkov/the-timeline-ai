#!/usr/bin/env python3
"""Isolated DOCX text extraction for Timeline document-extract sandboxes."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path


def extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                parts.append(line)
    return "\n".join(parts).strip()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    args = parser.parse_args()

    try:
        text = extract_docx(Path(args.input))
        json.dump({"ok": True, "text": text, "error": None}, sys.stdout)
        sys.stdout.write("\n")
        return 0
    except Exception as exc:  # noqa: BLE001
        json.dump({"ok": False, "text": "", "error": str(exc)}, sys.stdout)
        sys.stdout.write("\n")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
